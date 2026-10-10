"""Official household availability; never infer contracts, providers or geometry."""
import asyncio
import io
import math
import re
from datetime import UTC, datetime

from openpyxl import load_workbook
from publications import acquire, publish


def parse_workbook(body, municipalities):
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = book['Privathaushalte']
        rows = sheet.iter_rows(values_only=True)
        head = [next(rows) for _ in range(4)]
        match = re.fullmatch(r'Festnetz: (\d{2})\.(\d{4})', str(head[2][2]))
        if not match:
            raise ValueError('Missing fixed-network reference month')
        month, year = map(int, match.groups())
        observed = datetime(year, month, 1, tzinfo=UTC)
        if (head[3][:3] != ('AGS', 'Name', 'Verwaltungsebene')
                or head[2][6] != 'alle Technologien' or head[2][13] != 'FTTB/H'
                or list(head[3][6:14]) != [f'≥ {n} Mbit/s' for n in [16, 30, 50, 100, 200, 400, 1000, 1000]]
                or head[0][6] != 'Angaben zur Verfügbarkeit in Prozent der Privathaushalte'):
            raise ValueError('Unexpected broadband units or columns')
        areas = {}
        for row in rows:
            ags = str(row[0])
            if ags not in municipalities:
                continue
            if ags in areas or row[2] != '4 - Gemeinde':
                raise ValueError('Duplicate or invalid municipality')
            expected = municipalities[ags]
            if row[1] not in ('Gemeinde ' + expected, 'Stadt ' + expected):
                raise ValueError('Municipality code/name mismatch')
            values = []
            for raw in row[6:14]:
                if raw is None or raw in ('-', '–', '—'):
                    values.append(None)
                elif isinstance(raw, bool) or not isinstance(raw, (int, float)) or not math.isfinite(raw) or not 0 <= raw <= 100:
                    raise ValueError('Invalid household availability percentage')
                else:
                    values.append(raw)
            if all(v is None for v in values):
                raise ValueError('Municipality has no availability observations')
            areas[ags] = {'id': 'bba-' + ags, 'ags': ags, 'municipality': expected,
                          'referenceMonth': observed.strftime('%Y-%m'),
                          'fttbHPct': values[7], 'gigabitPct': values[6],
                          'availability': [{'minimumDownloadMbps': n, 'technology': 'all', 'percent': v}
                                           for n, v in zip([16, 30, 50, 100, 200, 400, 1000], values[:7])]}
        if set(areas) != set(municipalities):
            raise ValueError('Incomplete municipality coverage')
        return observed, {'contract_version': 'bba-households-v1', 'basis': 'published_statistics',
                          'referenceMonth': observed.strftime('%Y-%m'),
                          'unit': 'percent_of_private_households', 'publisher': 'Bundesnetzagentur',
                          'attribution': 'Breitbandatlas | Gigabit-Grundbuch (https://gigabitgrundbuch.bund.de)',
                          'selection': 'Auswahl der vier Gemeinden; unveränderte veröffentlichte Prozentwerte.',
                          'notice': 'Technisch verfügbare Anschlüsse, keine gemessene Geschwindigkeit oder Vertragsquote. FTTB/H umfasst Gebäude- und Wohnungsanschlüsse.',
                          'areas': list(areas.values())}
    finally:
        book.close()


async def import_broadband(conn, client, source):
    response, digest, attempt = await acquire(conn, client, source)
    observed, data = await asyncio.to_thread(parse_workbook, response.content, source['municipalities'])
    data['source_url'] = source['url']
    async with conn.transaction():
        await publish(conn, source, 'infrastructure/broadband', data, digest, observed)
        await conn.execute("UPDATE collection_attempts SET status='success',item_count=%s,item_count_unit='items' WHERE id=%s", (len(data["areas"]), attempt))
    await conn.commit()
