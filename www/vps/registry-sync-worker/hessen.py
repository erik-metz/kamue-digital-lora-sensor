"""Official HSL municipal workbook, preserving tables, periods, units and suppression.

These are published statistics, not live measurements or budget plans. The
workbook is discovered daily; no numbers are embedded in this adapter.
"""

import asyncio
import io
import re
from datetime import UTC, datetime
from html import unescape
from urllib.parse import urljoin, urlsplit

import openpyxl
from publications import acquire, publish

TABLES = {
    "demographics": ["1", "2", "3", "4"],
    "realestate": ["11", "12", "13"],
    "finance": ["16", "17", "18"],
    "economy": ["9", "14", "20", "21"],
    "environment": ["5", "6", "7", "8"],
    "social": ["15", "19"],
}


def clean(value):
    return re.sub(r"\s+", " ", str(value).replace("-\n", "")).strip()


def parse_workbook(body, municipalities):
    book = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    edition = str(book["Hessische Gemeindestatistik"]["A1"].value)
    if not re.fullmatch(r"Ausgabe 20\d{2}", edition):
        raise ValueError("Unexpected HSL publication identity")
    # Publication month comes from the publisher, not workbook filesystem metadata.
    imprint = " ".join(str(c.value or "") for row in book["Impressum"] for c in row)
    match = re.search(r"Erschienen im (\w+) (20\d{2})", imprint)
    months = [
        "Januar",
        "Februar",
        "März",
        "April",
        "Mai",
        "Juni",
        "Juli",
        "August",
        "September",
        "Oktober",
        "November",
        "Dezember",
    ]
    if not match or match[1] not in months:
        raise ValueError("Publication date missing")
    published = datetime(int(match[2]), months.index(match[1]) + 1, 1, tzinfo=UTC)
    result = {}
    for domain, numbers in TABLES.items():
        tables = []
        for number in numbers:
            sheet = book[number]
            first = next(
                (r for r in range(2, 20) if str(sheet.cell(r, 1).value) == "000000"),
                None,
            )
            if first is None:
                raise ValueError(f"HSL table {number}: missing region header")
            # Expand merged header cells only; never forward-fill missing observations.
            headers = {
                (r, c): sheet.cell(r, c).value
                for r in range(2, first)
                for c in range(3, sheet.max_column + 1)
            }
            for area in sheet.merged_cells.ranges:
                if area.min_row < 2 or area.max_row >= first:
                    continue
                value = sheet.cell(area.min_row, area.min_col).value
                for r in range(area.min_row, area.max_row + 1):
                    for c in range(area.min_col, area.max_col + 1):
                        headers[(r, c)] = value
            columns = []
            for col in range(3, sheet.max_column + 1):
                labels = list(
                    dict.fromkeys(
                        clean(headers[(r, col)])
                        for r in range(2, first)
                        if headers.get((r, col)) is not None
                    )
                )
                if not labels or "Zurück zum Inhalt" in labels:
                    continue
                columns.append({"column": col, "label": " / ".join(labels)})
            records = []
            for row in sheet.iter_rows(min_row=first):
                code = str(row[0].value)
                if code not in municipalities:
                    continue
                values = []
                for column in columns:
                    cell = row[column["column"] - 1]
                    raw = cell.value
                    value = (
                        raw
                        if isinstance(raw, (int, float))
                        else 0
                        if raw == "—"
                        else None
                    )
                    values.append(
                        {
                            "label": column["label"],
                            "value": value,
                            "source_marker": str(raw)
                            if value is None and raw is not None
                            else None,
                            "cell": cell.coordinate,
                        }
                    )
                records.append(
                    {
                        "municipality_id": municipalities[code],
                        "ags": "06" + code,
                        "name": row[1].value,
                        "values": values,
                    }
                )
            if {r["municipality_id"] for r in records} != set(municipalities.values()):
                raise ValueError(
                    f"HSL table {number}: incomplete municipality coverage"
                )
            tables.append(
                {"id": number, "title": clean(sheet["A1"].value), "records": records}
            )
        result["statistics/" + domain] = {
            "edition": edition,
            "publication_month": published.strftime("%Y-%m"),
            "basis": "published_statistics",
            "tables": tables,
        }
    book.close()
    return published, result


async def import_hessen(conn, client, source):
    listing, _, _ = await acquire(conn, client, source)
    links = re.findall(r"href=[\"\']([^\"\']+hgst_j\d{4}\.xlsx)", listing.text)
    urls = {urljoin(source["url"], unescape(link)) for link in links}
    if len(urls) != 1:
        raise ValueError("Cannot identify current official municipal workbook")
    url = urls.pop()
    if urlsplit(url).hostname != "statistik.hessen.de":
        raise ValueError("Unexpected workbook host")
    response, digest, attempt = await acquire(conn, client, source, url)
    published, datasets = await asyncio.to_thread(
        parse_workbook, response.content, source["municipalities"]
    )
    async with conn.transaction():
        for key, data in datasets.items():
            data["source_url"] = url
            await publish(conn, source, key, data, digest, published)
        await conn.execute(
            "UPDATE collection_attempts SET status='success' WHERE id=%s", (attempt,)
        )
    await conn.commit()
