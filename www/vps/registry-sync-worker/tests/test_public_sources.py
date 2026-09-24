import io
import unittest
from datetime import UTC, datetime

import openpyxl
from google.transit import gtfs_realtime_pb2
from hessen import parse_workbook, TABLES
from map_tiles import tile_inventory
from realtime import reported_delay


class PublicSourceTests(unittest.TestCase):
    def test_stop_delay_is_explicit_and_missing_delay_is_not_observed_on_time(self):
        update = gtfs_realtime_pb2.TripUpdate()
        stop = update.stop_time_update.add(stop_sequence=2)
        stop.departure.delay = 120
        metadata = {'stop_times':[{'stop_id':'s','sequence':2,'arrival':1000,'departure':1010}]}
        self.assertEqual(reported_delay(update,metadata,1000),(120,'next_reported_stop_approximation'))
        stop.schedule_relationship = 2
        self.assertEqual(reported_delay(update,metadata,1000),(0,'schedule_only'))
        update.delay = -30
        self.assertEqual(reported_delay(update,metadata,1000),(-30,'trip'))

    def test_tile_collection_is_bounded(self):
        tiles=tile_inventory([49.55,8.30,49.80,8.65],8,14)
        self.assertLess(len(tiles),3000)
        self.assertEqual(len(set(tiles)),len(tiles))
        with self.assertRaises(ValueError):
            tile_inventory([-80,-170,80,170],8,14)

    def test_workbook_preserves_period_units_and_suppression(self):
        book=openpyxl.Workbook()
        book.active.title='Hessische Gemeindestatistik'
        book.active['A1']='Ausgabe 2025'
        book.create_sheet('Impressum')['A1']='Erschienen im November 2025'
        for name in {n for names in TABLES.values() for n in names}:
            sheet=book.create_sheet(name)
            sheet.append(['Teststatistik 2020'])
            sheet.append(['Schlüssel','Gebiet','Fläche in Hektar','Unterdrückt','Null'])
            sheet.append(['000000','Land Hessen',100,None,None])
            sheet.append(['431005','Teststadt',1.5,'•','—'])
        buffer=io.BytesIO();book.save(buffer)
        date,data=parse_workbook(buffer.getvalue(),{'431005':'teststadt'})
        self.assertEqual(date,datetime(2025,11,1,tzinfo=UTC))
        table=data['statistics/environment']['tables'][0]
        self.assertEqual(table['title'],'Teststatistik 2020')
        values=table['records'][0]['values']
        self.assertEqual([r['value'] for r in values],[1.5,None,0])
        self.assertEqual(values[0]['label'],'Fläche in Hektar')
        self.assertEqual(values[1]['source_marker'],'•')
