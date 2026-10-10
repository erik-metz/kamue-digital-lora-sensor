import io
import unittest

from broadband import parse_workbook
from openpyxl import Workbook

MUNICIPALITIES = {'06431003': 'Biblis', '06431005': 'Bürstadt', '06431010': 'Groß-Rohrheim', '06431013': 'Lampertheim'}


def workbook(change=None):
    w = Workbook()
    s = w.active
    s.title = 'Privathaushalte'
    s.cell(1, 7, 'Angaben zur Verfügbarkeit in Prozent der Privathaushalte')
    s.cell(3, 3, 'Festnetz: 12.2025')
    s.cell(3, 7, 'alle Technologien')
    s.cell(3, 14, 'FTTB/H')
    for i, v in enumerate(['AGS', 'Name', 'Verwaltungsebene'], 1): s.cell(4, i, v)
    for i, v in enumerate([16, 30, 50, 100, 200, 400, 1000, 1000], 7): s.cell(4, i, f'≥ {v} Mbit/s')
    for row, (ags, name) in enumerate(MUNICIPALITIES.items(), 5):
        s.cell(row, 1, ags); s.cell(row, 2, 'Gemeinde ' + name); s.cell(row, 3, '4 - Gemeinde')
        for col in range(7, 15): s.cell(row, col, 50.25)
    if change: change(s)
    out = io.BytesIO(); w.save(out); return out.getvalue()


class BroadbandTests(unittest.TestCase):
    def test_original_percentages_and_missing_are_distinct_from_zero(self):
        def change(s):
            s['N5'] = 0
            s['N6'] = '—'
        stamp, data = parse_workbook(workbook(change), MUNICIPALITIES)
        self.assertEqual(stamp.isoformat(), '2025-12-01T00:00:00+00:00')
        self.assertEqual(len(data['areas']), 4)
        self.assertEqual(data['areas'][0]['fttbHPct'], 0)
        self.assertIsNone(data['areas'][1]['fttbHPct'])
        self.assertEqual(data['areas'][0]['gigabitPct'], 50.25)

    def test_reject_bad_units_reference_missing_duplicate_and_invalid_values(self):
        for cell, value in [('G1', 'percent_of_area'), ('N3', 'FTTH'), ('C3', 'Festnetz: 13.2025'),
                            ('A5', '06431099'), ('A6', '06431003'), ('B5', 'Gemeinde Falsch'),
                            ('N5', 101), ('N5', True), ('N5', 'unbekannt')]:
            with self.subTest(cell=cell, value=value), self.assertRaises((ValueError, KeyError)):
                parse_workbook(workbook(lambda s, cell=cell, value=value: setattr(s[cell], 'value', value)), MUNICIPALITIES)
